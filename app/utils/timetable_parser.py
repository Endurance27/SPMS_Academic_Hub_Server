import re
from typing import List, Dict, Any, Optional

import openpyxl
import csv
import io


class TimetableParseIssue:
    def __init__(self, row: int, severity: str, message: str):
        self.row = row
        self.severity = severity
        self.message = message


# Day columns in the pivot layout: (subject_col_index, venue_col_index, day_index)
# Columns are 0-based: B=1, C=2, D=3, E=4, ...
_DAY_COLUMNS = [
    (1, 2, 0),   # Monday:    B/C
    (3, 4, 1),   # Tuesday:   D/E
    (5, 6, 2),   # Wednesday: F/G
    (7, 8, 3),   # Thursday:  H/I
    (9, 10, 4),  # Friday:    J/K
]

_WEEKDAY_NAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

# Map start time "HH:MM" → period index (0-based)
_TIME_TO_PERIOD = {}
for _i in range(12):
    _h = 7 + _i
    _TIME_TO_PERIOD[f"{_h}:30"] = _i
    _TIME_TO_PERIOD[f"{_h:02d}:30"] = _i


def _parse_time_slot(raw: str) -> Optional[int]:
    """Extract period index from a time-slot string like '7:30 - 8:20'."""
    if not raw:
        return None
    cleaned = re.sub(r"\s+", "", raw)
    m = re.match(r"(\d{1,2}):(\d{2})", cleaned)
    if not m:
        return None
    key = f"{int(m.group(1))}:{m.group(2)}"
    return _TIME_TO_PERIOD.get(key)


def _extract_level_from_sheet(name: str) -> Optional[int]:
    """Extract numeric level from sheet name like 'L100v1'."""
    m = re.search(r"L?(\d{3})", name, re.IGNORECASE)
    return int(m.group(1)) if m else None


def _normalize_subject(subject: str) -> str:
    subject = subject.upper().strip()
    subject = re.sub(r"\s+", " ", subject)
    subject = subject.replace("©", "(C)")
    subject = re.sub(r"\(\s*C\s*\)", "(C)", subject)
    subject = re.sub(r"\(\s*E\s*\)", "(E)", subject)
    return subject


def _normalize_venue(venue: str) -> str:
    venue = venue.upper().strip()
    venue = re.sub(r"\s+", " ", venue)
    venue = re.sub(r"^L/OT\s*", "LOT", venue)
    venue = re.sub(r"^LOT\s+", "LOT", venue)
    venue = re.sub(r"^JQB\s*(\d+)$", r"JQB\1", venue)
    venue = re.sub(r"^GCB\s*FFR(\d)", r"GCB FF R\1", venue)
    venue = re.sub(r"^GCB\s*FF\s*RM(\d)", r"GCB FF R\1", venue)
    return venue


def _extract_course_kind(subject: str) -> Optional[str]:
    if "(C)" in subject:
        return "C"
    if "(E)" in subject:
        return "E"
    return None


def _infer_level_from_subject(subject: str) -> Optional[int]:
    """Infer academic level from a course code like 'MATH 353' → 300."""
    m = re.search(r"[A-Z]+\s*(\d)\d{2}", subject)
    if m:
        return int(m.group(1)) * 100
    return None


def _strip_lecturer_from_venue(raw: str) -> str:
    """Strip trailing parenthesised lecturer initials from venue, e.g. 'MATH05 (BFS)' → 'MATH05'."""
    return re.sub(r"\s*\([A-Za-z0-9]{1,6}\)\s*$", "", raw.strip())


def _sheet_has_timetable_content(ws) -> bool:
    """Check whether a worksheet looks like a timetable pivot (has day headers and time rows)."""
    for row in ws.iter_rows(min_row=1, max_row=min(10, ws.max_row), values_only=True):
        text = " ".join(str(c).upper() for c in row if c is not None)
        if any(d in text for d in ("MONDAY", "TUESDAY", "WEDNESDAY")):
            return True
    return False


class TimetableParser:
    """Parser for University of Ghana master timetable Excel files.

    Handles the pivot-table format where:
    - Each sheet is one level (L100v1, L200v1, ...)
    - Days are paired columns (Subject/Venue) across the row
    - Time slots span multiple data rows vertically
    - The layout repeats with section headers every ~25 rows
    """

    def __init__(self):
        self.issues: List[TimetableParseIssue] = []
        self.entries: List[Dict[str, Any]] = []

    def parse_file(
        self,
        file_content: bytes,
        file_type: str,
        level: int,
        semester_id: str,
    ) -> Dict[str, Any]:
        """Parse uploaded timetable file.

        Args:
            file_content: Raw bytes of the uploaded file.
            file_type: "xlsx" or "csv".
            level: Target level (100/200/300/400). 0 means all levels.
            semester_id: Semester UUID string (unused by parser, kept for API compat).
        """
        self.issues = []
        self.entries = []

        try:
            if file_type == "xlsx":
                self._parse_excel(file_content, level)
            elif file_type == "csv":
                self._parse_csv(file_content, level)
            else:
                self.issues.append(
                    TimetableParseIssue(0, "error", "Upload the master timetable as an Excel (.xlsx) or CSV file")
                )
        except Exception as e:
            self.issues.append(TimetableParseIssue(0, "error", f"File parsing error: {str(e)}"))

        if not self.entries and not self.issues:
            self.issues.append(TimetableParseIssue(0, "error", "No timetable entries found in the file"))

        free_count = sum(1 for e in self.entries if e.get("is_free"))
        venues = set(e["venue_code"] for e in self.entries)
        levels_found = sorted(set(e["level"] for e in self.entries))

        return {
            "entries": self.entries,
            "issues": [
                {"row": i.row, "severity": i.severity, "message": i.message}
                for i in self.issues
            ],
            "summary": {
                "entry_count": len(self.entries),
                "free_count": free_count,
                "venue_count": len(venues),
                "levels": levels_found,
                "level": level if level else None,
            },
        }

    # ------------------------------------------------------------------
    # Excel pivot-table parser (primary format)
    # ------------------------------------------------------------------

    def _parse_excel(self, content: bytes, level: int):
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

        sheets_parsed = 0
        for sheet_name in wb.sheetnames:
            sheet_level = _extract_level_from_sheet(sheet_name)
            if sheet_level is None:
                continue
            if level and sheet_level != level:
                continue

            ws = wb[sheet_name]
            self._parse_pivot_sheet(ws, sheet_level, sheet_name)
            sheets_parsed += 1

        # Fallback: no level-named sheets found — try sheets with timetable content
        # (common for department timetable files that use a single "Timetable" sheet)
        if sheets_parsed == 0:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                if _sheet_has_timetable_content(ws):
                    self._parse_pivot_sheet(ws, 0, sheet_name, infer_level=True)
                    sheets_parsed += 1
                    break

        if sheets_parsed == 0:
            if level:
                available = [s for s in wb.sheetnames if _extract_level_from_sheet(s)]
                self.issues.append(
                    TimetableParseIssue(
                        0, "error",
                        f"No sheet found for Level {level}. "
                        f"Available sheets: {', '.join(available) if available else 'none'}"
                    )
                )
            else:
                self.issues.append(TimetableParseIssue(0, "error", "No level sheets found in the workbook"))

    def _parse_pivot_sheet(self, ws, level: int, sheet_name: str, infer_level: bool = False):
        """Parse one sheet in the UG master-timetable pivot format."""
        current_period: Optional[int] = None
        row_idx = 0

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
            row_idx += 1
            col_a = str(row[0]).strip() if row[0] is not None else ""

            # Skip level-label rows (e.g. "L100") and header rows ("Time")
            if re.match(r"^L?\d{3}$", col_a, re.IGNORECASE):
                continue
            if col_a.lower() == "time":
                continue

            # Stop at non-timetable sections (lecturer lists, metadata)
            if re.match(r"(?i)^LECTURER|^COURSE\s", col_a):
                break

            # Check if column A has a new time slot
            if col_a:
                period = _parse_time_slot(col_a)
                if period is not None:
                    current_period = period

            if current_period is None:
                continue

            # Extract subject/venue pairs for each day
            for subj_col, venue_col, day_idx in _DAY_COLUMNS:
                subj_raw = row[subj_col] if subj_col < len(row) else None
                venue_raw = row[venue_col] if venue_col < len(row) else None

                if subj_raw is None or str(subj_raw).strip() == "":
                    continue

                subject = _normalize_subject(str(subj_raw))

                # Skip non-course entries (lecturer names, labels, level headers)
                if infer_level and not re.match(r"^[A-Z]{2,6}\s*\d{3}", subject):
                    continue

                # Clean venue: strip lecturer initials in parentheses, then normalise
                venue_text = str(venue_raw).strip() if venue_raw and str(venue_raw).strip() else ""
                if venue_text:
                    venue_text = _strip_lecturer_from_venue(venue_text)
                venue = _normalize_venue(venue_text) if venue_text else ""

                is_free = subject == "FREE"

                if not venue and not is_free:
                    self.issues.append(
                        TimetableParseIssue(
                            row_idx, "warning",
                            f"[{sheet_name}] {_WEEKDAY_NAMES[day_idx]} period {current_period}: "
                            f"subject '{subject}' has no venue"
                        )
                    )

                entry_level = level
                if infer_level:
                    entry_level = _infer_level_from_subject(subject) or 0

                self.entries.append({
                    "level": entry_level,
                    "day": day_idx,
                    "period": current_period,
                    "subject": subject,
                    "venue_code": venue,
                    "is_free": is_free,
                    "course_kind": _extract_course_kind(subject) if not is_free else None,
                })

    # ------------------------------------------------------------------
    # CSV fallback (flat format: Level, Day, Time, Subject, Venue)
    # ------------------------------------------------------------------

    def _parse_csv(self, content: bytes, level: int):
        text = content.decode("utf-8")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
        if not rows:
            self.issues.append(TimetableParseIssue(0, "error", "The file is empty"))
            return
        self._parse_flat_rows(rows, level)

    def _parse_flat_rows(self, rows: List, level: int):
        header = [str(c).lower().strip() for c in rows[0]]

        required = {"day", "subject", "venue"}
        present = set(header)
        missing = required - present
        if missing:
            self.issues.append(
                TimetableParseIssue(
                    1, "error",
                    f"Missing column(s): {', '.join(sorted(missing))}. Expected: Day, Time, Subject, Venue"
                )
            )
            return

        col = {name: header.index(name) for name in header if name}
        has_time = "time" in col or "start" in col

        weekday_map = {d[:3].lower(): i for i, d in enumerate(_WEEKDAY_NAMES)}

        for row_num, row in enumerate(rows[1:], start=2):
            if not any(str(c).strip() for c in row):
                continue

            day_raw = row[col["day"]].strip() if "day" in col and col["day"] < len(row) else ""
            day_idx = weekday_map.get(day_raw[:3].lower()) if day_raw else None
            if day_idx is None:
                self.issues.append(TimetableParseIssue(row_num, "error", f"Unknown day: {day_raw}"))
                continue

            time_raw = ""
            if "time" in col and col["time"] < len(row):
                time_raw = row[col["time"]].strip()
            elif "start" in col and col["start"] < len(row):
                time_raw = row[col["start"]].strip()
            period = _parse_time_slot(time_raw)
            if period is None:
                self.issues.append(TimetableParseIssue(row_num, "error", f"Invalid time: {time_raw}"))
                continue

            subj_raw = row[col["subject"]].strip() if "subject" in col and col["subject"] < len(row) else ""
            if not subj_raw:
                self.issues.append(TimetableParseIssue(row_num, "error", "Empty subject"))
                continue

            venue_raw = row[col["venue"]].strip() if "venue" in col and col["venue"] < len(row) else ""
            if not venue_raw:
                self.issues.append(TimetableParseIssue(row_num, "error", "Empty venue"))
                continue

            subject = _normalize_subject(subj_raw)
            venue = _normalize_venue(venue_raw)
            is_free = subject == "FREE"

            lvl_raw = row[col["level"]].strip() if "level" in col and col["level"] < len(row) else ""
            row_level = level
            if lvl_raw:
                m = re.search(r"\d{3}", lvl_raw)
                if m:
                    row_level = int(m.group())

            if level and row_level != level:
                continue

            self.entries.append({
                "level": row_level,
                "day": day_idx,
                "period": period,
                "subject": subject,
                "venue_code": venue,
                "is_free": is_free,
                "course_kind": _extract_course_kind(subject) if not is_free else None,
            })
